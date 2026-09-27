# -*- coding: utf-8 -*-
import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_paragraph_rtl(p):
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1')
    pPr.append(bidi)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

def set_table_rtl(table):
    tblPr = table._tbl.tblPr
    bidiVisual = OxmlElement('w:bidiVisual')
    tblPr.append(bidiVisual)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    set_paragraph_rtl(h)
    run = h.runs[0] if h.runs else h.add_run()
    run.font.name = 'Segoe UI'
    if level == 1:
        run.font.size = Pt(20)
        run.font.color.rgb = RGBColor(14, 116, 144) # Deep Cyan
        run.bold = True
    elif level == 2:
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(30, 64, 175) # Royal Blue
        run.bold = True
    elif level == 3:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(107, 33, 168) # Deep Purple
        run.bold = True
    return h

def main():
    guide_path = "USER_GUIDE.md"
    docx_path = "USER_GUIDE.docx"

    if not os.path.exists(guide_path):
        print(f"Error: {guide_path} not found")
        return

    with open(guide_path, "r", encoding="utf-8") as f:
        content = f.read()

    doc = Document()

    # הגדרות עמוד
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # כותרת ראשית
    title_p = doc.add_paragraph()
    set_paragraph_rtl(title_p)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("📘 מדריך למשתמש – פייתון בקלות")
    title_run.font.name = "Segoe UI"
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42)

    sub_p = doc.add_paragraph()
    set_paragraph_rtl(sub_p)
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("המדריך המלא לשימוש באפליקציה, ניצול כל הפיצ'רים וטיפים מעשיים להצלחה והתמדה")
    sub_run.font.name = "Segoe UI"
    sub_run.font.size = Pt(13)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph() # מרווח

    lines = content.split("\n")
    in_table = False
    table_lines = []

    def flush_table(lines_to_parse):
        if not lines_to_parse:
            return
        rows_data = []
        for l in lines_to_parse:
            if not l.strip() or "---" in l:
                continue
            cells = [c.strip() for c in l.strip().strip("|").split("|")]
            if len(cells) >= 2:
                rows_data.append(cells)
        
        if not rows_data:
            return

        cols = len(rows_data[0])
        table = doc.add_table(rows=len(rows_data), cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_rtl(table)

        for i, row in enumerate(rows_data):
            for j, cell_text in enumerate(row):
                cell = table.cell(i, j)
                cell.text = cell_text.replace("**", "").replace("\\_", "_")
                p = cell.paragraphs[0]
                set_paragraph_rtl(p)
                p.runs[0].font.name = "Segoe UI"
                p.runs[0].font.size = Pt(10.5)

                if i == 0:
                    # כותרת טבלה
                    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0F172A"/>')
                    cell._tc.get_or_add_tcPr().append(shd)
                    p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
                    p.runs[0].font.bold = True
                else:
                    if i % 2 == 1:
                        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
                        cell._tc.get_or_add_tcPr().append(shd)
                    p.runs[0].font.color.rgb = RGBColor(30, 41, 59)

        doc.add_paragraph()

    i = 0
    while i < len(lines):
        line = lines[i]
        trimmed = line.strip()

        # התעלמות מכותרת ראשית שעיצבנו כבר
        if trimmed.startswith("# 📘 מדריך למשתמש") or trimmed.startswith("> **ברוך הבא"):
            i += 1
            continue

        if trimmed.startswith("|") and "|" in trimmed[1:]:
            table_lines.append(trimmed)
            i += 1
            continue
        elif table_lines:
            flush_table(table_lines)
            table_lines = []

        if not trimmed:
            i += 1
            continue

        if trimmed == "---":
            # קו מפריד
            i += 1
            continue

        if trimmed.startswith("## "):
            text = trimmed[3:].strip()
            add_styled_heading(doc, text, level=1)
        elif trimmed.startswith("### "):
            text = trimmed[4:].strip()
            add_styled_heading(doc, text, level=2)
        elif trimmed.startswith("#### "):
            text = trimmed[5:].strip()
            add_styled_heading(doc, text, level=3)
        elif trimmed.startswith("* ") or trimmed.startswith("- "):
            bullet_text = trimmed[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            set_paragraph_rtl(p)
            parts = re.split(r'(\*\*.*?\*\*)', bullet_text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                    r.font.name = "Segoe UI"
                    r.font.color.rgb = RGBColor(30, 41, 59)
                else:
                    r = p.add_run(part)
                    r.font.name = "Segoe UI"
                    r.font.color.rgb = RGBColor(51, 65, 85)
        elif re.match(r'^\d+\.\s', trimmed):
            # רשימה ממוספרת
            num_text = re.sub(r'^\d+\.\s', '', trimmed)
            p = doc.add_paragraph(style='List Number')
            set_paragraph_rtl(p)
            parts = re.split(r'(\*\*.*?\*\*)', num_text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                    r.font.name = "Segoe UI"
                    r.font.color.rgb = RGBColor(30, 41, 59)
                else:
                    r = p.add_run(part)
                    r.font.name = "Segoe UI"
                    r.font.color.rgb = RGBColor(51, 65, 85)
        else:
            p = doc.add_paragraph()
            set_paragraph_rtl(p)
            parts = re.split(r'(\*\*.*?\*\*)', trimmed)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                    r.font.name = "Segoe UI"
                    r.font.color.rgb = RGBColor(15, 23, 42)
                else:
                    r = p.add_run(part)
                    r.font.name = "Segoe UI"
                    r.font.color.rgb = RGBColor(51, 65, 85)

        i += 1

    if table_lines:
        flush_table(table_lines)

    doc.save(docx_path)
    print(f"Successfully converted to {docx_path}")

if __name__ == "__main__":
    main()
